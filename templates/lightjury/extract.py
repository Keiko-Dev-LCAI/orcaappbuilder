"""LightJury — local document text extraction for uploaded evidence.

Pulls plain text out of a PDF, Word (.docx), or text file ON THE SERVER — no
third-party API, so a party's evidence never leaves the pipeline. Images are out
of scope for v1 (the jury model is text-only); callers should reject them.
"""

import io
import os

MAX_BYTES = 8 * 1024 * 1024          # 8 MB per file
MAX_TEXT_CHARS = 60_000              # cap extracted text so one file can't flood
ALLOWED_EXT = {".pdf", ".docx", ".txt", ".md", ".rtf", ".csv"}


class ExtractError(Exception):
    pass


def allowed(filename):
    return os.path.splitext((filename or "").lower())[1] in ALLOWED_EXT


def _from_txt(data):
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return data.decode("utf-8", errors="replace")


def _from_pdf(data):
    try:
        from pypdf import PdfReader
    except Exception as e:
        raise ExtractError(f"PDF support unavailable: {e}")
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as e:
        raise ExtractError(f"Could not read PDF: {str(e)[:120]}")
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            raise ExtractError("This PDF is password-protected.")
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _from_docx(data):
    try:
        import docx  # python-docx
    except Exception as e:
        raise ExtractError(f"Word support unavailable: {e}")
    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractError(f"Could not read Word file: {str(e)[:120]}")
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lines)


def extract_text(filename, data):
    """Return cleaned text extracted from a file's bytes.

    Raises ExtractError with a user-safe message on any problem.
    """
    if not filename:
        raise ExtractError("Missing filename.")
    if not allowed(filename):
        raise ExtractError("Unsupported file type. Use PDF, Word (.docx), or text.")
    if data is None or len(data) == 0:
        raise ExtractError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ExtractError("File is too large (8 MB max).")

    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        text = _from_pdf(data)
    elif ext == ".docx":
        text = _from_docx(data)
    else:  # .txt / .md / .rtf / .csv
        text = _from_txt(data)

    text = (text or "").strip()
    if not text:
        raise ExtractError(
            "No readable text found. If this is a scan or photo, it needs "
            "image analysis, which isn't supported yet.")
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n[...truncated...]"
    return text
